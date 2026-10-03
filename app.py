import os
import re
import time
import smtplib
import ssl
from datetime import datetime
from email.message import EmailMessage
from functools import wraps

from dotenv import load_dotenv
from flask import Flask, render_template, request, redirect, url_for, flash, session, send_file
from flask_sqlalchemy import SQLAlchemy
from flask_login import LoginManager, UserMixin, login_user, login_required, logout_user, current_user
from openpyxl import load_workbook
from werkzeug.security import generate_password_hash, check_password_hash
from email_validator import validate_email, EmailNotValidError
from sqlalchemy import or_

load_dotenv()

db = SQLAlchemy()
login_manager = LoginManager()

class User(UserMixin, db.Model):
    id = db.Column(db.Integer, primary_key=True)
    username = db.Column(db.String(120), unique=True, nullable=False)
    password_hash = db.Column(db.String(255), nullable=False)

    def set_password(self, p):
        self.password_hash = generate_password_hash(p)

    def check_password(self, p):
        return check_password_hash(self.password_hash, p)

class Contact(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    company = db.Column(db.String(255))
    email = db.Column(db.String(320), nullable=False, index=True)
    website = db.Column(db.String(1000))
    contact_name = db.Column(db.String(255))
    business_type = db.Column(db.String(255))
    country = db.Column(db.String(120))
    phone = db.Column(db.String(100))
    status = db.Column(db.String(40), default="pending", index=True)
    last_error = db.Column(db.Text)
    sent_at = db.Column(db.DateTime)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)

class Campaign(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    subject = db.Column(db.String(500), nullable=False)
    body = db.Column(db.Text, nullable=False)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)

class Suppression(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    email = db.Column(db.String(320), unique=True, nullable=False, index=True)
    reason = db.Column(db.String(255))
    created_at = db.Column(db.DateTime, default=datetime.utcnow)

@login_manager.user_loader
def load_user(user_id):
    return db.session.get(User, int(user_id))

def create_app():
    app = Flask(__name__, instance_relative_config=True)
    app.config["SECRET_KEY"] = os.getenv("SECRET_KEY", "dev-change-me")
    app.config["SQLALCHEMY_DATABASE_URI"] = "sqlite:///" + os.path.join(app.instance_path, "outreach.db")
    app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False
    os.makedirs(app.instance_path, exist_ok=True)

    db.init_app(app)
    login_manager.init_app(app)
    login_manager.login_view = "login"

    with app.app_context():
        db.create_all()
        if not User.query.filter_by(username=os.getenv("ADMIN_USERNAME", "admin")).first():
            u = User(username=os.getenv("ADMIN_USERNAME", "admin"))
            u.set_password(os.getenv("ADMIN_PASSWORD", "change-this-password"))
            db.session.add(u)
            db.session.commit()

    @app.route("/")
    @login_required
    def dashboard():
        total = Contact.query.count()
        pending = Contact.query.filter_by(status="pending").count()
        sent = Contact.query.filter_by(status="sent").count()
        failed = Contact.query.filter_by(status="failed").count()
        suppressed = Suppression.query.count()
        campaigns = Campaign.query.order_by(Campaign.created_at.desc()).limit(5).all()
        return render_template("dashboard.html", total=total, pending=pending,
                               sent=sent, failed=failed, suppressed=suppressed,
                               campaigns=campaigns)

    @app.route("/campaigns")
    @login_required
    def campaigns():
        rows = Campaign.query.order_by(Campaign.created_at.desc()).all()
        return render_template("campaigns.html", campaigns=rows,
                               pending=Contact.query.filter_by(status="pending").count())

    @app.route("/campaigns/<int:campaign_id>/duplicate", methods=["POST"])
    @login_required
    def duplicate_campaign(campaign_id):
        source = db.session.get(Campaign, campaign_id)
        if not source:
            flash("Campaign not found.", "error")
            return redirect(url_for("campaigns"))
        duplicate = Campaign(subject=f"{source.subject} (copy)"[:500], body=source.body)
        db.session.add(duplicate)
        db.session.commit()
        flash("Campaign copied. Edit it before sending.", "success")
        return redirect(url_for("compose", campaign_id=duplicate.id))

    @app.route("/login", methods=["GET", "POST"])
    def login():
        if request.method == "POST":
            u = User.query.filter_by(username=request.form["username"].strip()).first()
            if u and u.check_password(request.form["password"]):
                login_user(u)
                return redirect(url_for("dashboard"))
            flash("Invalid username or password.", "error")
        return render_template("login.html")

    @app.route("/logout")
    @login_required
    def logout():
        logout_user()
        return redirect(url_for("login"))

    @app.route("/import", methods=["GET", "POST"])
    @login_required
    def import_excel():
        if request.method == "POST":
            f = request.files.get("file")
            if not f or not f.filename.lower().endswith((".xlsx", ".xlsm")):
                flash("Please upload an .xlsx file.", "error")
                return redirect(url_for("import_excel"))

            wb = load_workbook(f, read_only=True, data_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                flash("The spreadsheet is empty.", "error")
                return redirect(url_for("import_excel"))

            headers = [str(x).strip().lower() if x is not None else "" for x in rows[0]]
            def col(*names):
                for n in names:
                    if n in headers:
                        return headers.index(n)
                return None

            idx_email = col("emailid1", "email", "email_id")
            idx_url = col("url", "website", "web")
            idx_company = col("clientname", "company", "companyname", "company_name")
            idx_contact_name = col("contactname", "contact_name", "name", "firstname", "first_name")
            idx_type = col("bussinesstype", "business_type")
            idx_country = col("country")
            idx_phone = col("phoneormobile", "phone", "mobile")

            if idx_email is None:
                flash("No email column found. Expected emailid1 or email.", "error")
                return redirect(url_for("import_excel"))

            added = 0
            skipped = 0
            existing = {x.email.lower() for x in Contact.query.with_entities(Contact.email).all()}
            suppressed = {x.email.lower() for x in Suppression.query.with_entities(Suppression.email).all()}

            for row in rows[1:]:
                raw = row[idx_email] if idx_email < len(row) else None
                if not raw:
                    skipped += 1
                    continue
                email = str(raw).strip().lower()
                try:
                    email = validate_email(email, check_deliverability=False).normalized
                except EmailNotValidError:
                    skipped += 1
                    continue

                if email in existing or email in suppressed:
                    skipped += 1
                    continue

                def val(i):
                    return str(row[i]).strip() if i is not None and i < len(row) and row[i] is not None else ""

                db.session.add(Contact(
                    email=email,
                    website=val(idx_url),
                    company=val(idx_company),
                    contact_name=val(idx_contact_name),
                    business_type=val(idx_type),
                    country=val(idx_country),
                    phone=val(idx_phone),
                ))
                existing.add(email)
                added += 1

            db.session.commit()
            flash(f"Imported {added} contacts; skipped {skipped}.", "success")
            return redirect(url_for("dashboard"))

        return render_template("import.html")

    @app.route("/compose", methods=["GET", "POST"])
    @login_required
    def compose():
        campaign_id = request.values.get("campaign_id", type=int)
        campaign = db.session.get(Campaign, campaign_id) if campaign_id else None
        if campaign_id and not campaign:
            flash("Campaign not found.", "error")
            return redirect(url_for("campaigns"))
        if request.method == "POST":
            subject = request.form["subject"].strip()
            body = request.form["body"].strip()
            if not subject or not body:
                flash("Subject and body are required.", "error")
                return redirect(url_for("compose"))
            if campaign:
                campaign.subject = subject
                campaign.body = body
                c = campaign
            else:
                c = Campaign(subject=subject, body=body)
                db.session.add(c)
            db.session.commit()
            flash("Campaign updated." if campaign_id else "Campaign saved.", "success")
            return redirect(url_for("preview", campaign_id=c.id))
        return render_template("compose.html", campaign=campaign)

    def render_text(template, contact):
        values = {
            "company": contact.company or "",
            "clientname": contact.company or "",
            "contact_name": contact.contact_name or "",
            "email": contact.email,
            "url": contact.website or "",
            "website": contact.website or "",
            "bussinesstype": contact.business_type or "",
            "business_type": contact.business_type or "",
            "country": contact.country or "",
            "phone": contact.phone or "",
        }
        out = template
        for k, v in values.items():
            out = out.replace("{{" + k + "}}", v)
        return out

    @app.route("/preview/<int:campaign_id>")
    @login_required
    def preview(campaign_id):
        campaign = db.session.get(Campaign, campaign_id)
        sample = (Contact.query.filter_by(status="pending").first()
                  or Contact.query.order_by(Contact.id.desc()).first())
        try:
            batch_limit = max(1, min(int(os.getenv("MAX_SENDS_PER_RUN", "20")), 100))
        except ValueError:
            batch_limit = 20
        if not campaign or not sample:
            flash("Create a campaign and import at least one pending contact.", "error")
            return redirect(url_for("dashboard"))
        return render_template("preview.html", campaign=campaign, sample=sample,
                               subject=render_text(campaign.subject, sample),
                               body=render_text(campaign.body, sample),
                               pending=Contact.query.filter_by(status="pending").count(),
                               batch_limit=batch_limit)

    def send_email(to_addr, subject, body):
        host = os.getenv("SMTP_HOST")
        try:
            port = int(os.getenv("SMTP_PORT", "587"))
        except ValueError as exc:
            raise RuntimeError("SMTP_PORT must be a valid port number") from exc
        username = os.getenv("SMTP_USERNAME")
        password = os.getenv("SMTP_PASSWORD")
        use_tls = os.getenv("SMTP_USE_TLS", "true").lower() == "true"

        if not host or not username or not password:
            raise RuntimeError("SMTP_HOST, SMTP_USERNAME, and SMTP_PASSWORD must be set in .env")

        msg = EmailMessage()
        msg["From"] = username
        msg["To"] = to_addr
        msg["Subject"] = subject
        msg.set_content(body)

        if port == 465:
            smtp_connection = smtplib.SMTP_SSL(
                host, port, timeout=30, context=ssl.create_default_context()
            )
        else:
            smtp_connection = smtplib.SMTP(host, port, timeout=30)

        with smtp_connection as smtp:
            if use_tls and port != 465:
                smtp.starttls(context=ssl.create_default_context())
            smtp.login(username, password)
            smtp.send_message(msg)

    @app.route("/send", methods=["POST"])
    @login_required
    def send_campaign():
        campaign_id = int(request.form["campaign_id"])
        campaign = db.session.get(Campaign, campaign_id)
        if not campaign:
            flash("Campaign not found.", "error")
            return redirect(url_for("dashboard"))

        try:
            max_sends = max(1, min(int(os.getenv("MAX_SENDS_PER_RUN", "20")), 100))
            interval = max(1, int(os.getenv("SEND_INTERVAL_SECONDS", "10")))
        except ValueError:
            max_sends, interval = 20, 10

        contacts = Contact.query.filter_by(status="pending").limit(max_sends).all()
        if not contacts:
            flash("No pending contacts.", "error")
            return redirect(url_for("dashboard"))

        sent_count = 0
        failed_count = 0
        suppressed_count = 0
        first_error = None
        for i, contact in enumerate(contacts):
            if Suppression.query.filter_by(email=contact.email).first():
                contact.status = "suppressed"
                suppressed_count += 1
                db.session.commit()
                continue
            try:
                send_email(
                    contact.email,
                    render_text(campaign.subject, contact),
                    render_text(campaign.body, contact)
                )
                contact.status = "sent"
                contact.sent_at = datetime.utcnow()
                contact.last_error = None
                sent_count += 1
            except Exception as e:
                contact.status = "failed"
                contact.last_error = str(e)[:2000]
                failed_count += 1
                if first_error is None:
                    first_error = str(e).splitlines()[0][:240]
            db.session.commit()

            if i < len(contacts) - 1:
                time.sleep(interval)

        summary = f"Processed {len(contacts)} contacts; sent {sent_count}, failed {failed_count}, suppressed {suppressed_count}."
        if first_error:
            flash(f"{summary} First error: {first_error}", "error")
        else:
            flash(summary, "success")
        return redirect(url_for("dashboard"))

    @app.route("/contacts")
    @login_required
    def contacts():
        query = request.args.get("q", "").strip()
        status = request.args.get("status", "all")
        page = max(1, request.args.get("page", 1, type=int))
        per_page = 25
        contacts_query = Contact.query
        if query:
            pattern = f"%{query}%"
            contacts_query = contacts_query.filter(
                or_(Contact.email.ilike(pattern), Contact.company.ilike(pattern),
                    Contact.contact_name.ilike(pattern))
            )
        if status in {"pending", "sent", "failed", "suppressed"}:
            contacts_query = contacts_query.filter_by(status=status)
        total_rows = contacts_query.count()
        pages = max(1, (total_rows + per_page - 1) // per_page)
        page = min(page, pages)
        rows = contacts_query.order_by(Contact.id.desc()).offset((page - 1) * per_page).limit(per_page).all()
        return render_template("contacts.html", contacts=rows, q=query,
                               status=status, page=page, pages=pages,
                               total_rows=total_rows)

    @app.route("/contacts/<int:contact_id>/edit", methods=["GET", "POST"])
    @login_required
    def edit_contact(contact_id):
        contact = db.session.get(Contact, contact_id)
        if not contact:
            flash("Contact not found.", "error")
            return redirect(url_for("contacts"))
        fields = ("company", "contact_name", "website", "business_type", "country", "phone")
        if request.method == "POST":
            limits = {"company": 255, "contact_name": 255, "website": 1000,
                      "business_type": 255, "country": 120, "phone": 100}
            values = {field: request.form.get(field, "").strip() for field in fields}
            oversized = [field for field, value in values.items() if len(value) > limits[field]]
            if oversized:
                flash("One or more fields are too long. Shorten the values and save again.", "error")
                return render_template("edit_contact.html", contact=contact, values=values)
            for field, value in values.items():
                setattr(contact, field, value or None)
            db.session.commit()
            flash("Contact details updated.", "success")
            return redirect(url_for("contacts", q=contact.email))
        values = {field: getattr(contact, field) or "" for field in fields}
        return render_template("edit_contact.html", contact=contact, values=values)

    @app.route("/contacts/bulk-retry", methods=["POST"])
    @login_required
    def bulk_retry_contacts():
        raw_ids = request.form.getlist("contact_ids")
        try:
            contact_ids = {int(value) for value in raw_ids}
        except ValueError:
            flash("Invalid contact selection.", "error")
            return redirect(url_for("contacts", status="failed"))
        if not contact_ids:
            flash("Select at least one failed contact to queue for retry.", "error")
            return redirect(url_for("contacts", status="failed"))

        queued = 0
        for contact in Contact.query.filter(Contact.id.in_(contact_ids), Contact.status == "failed").all():
            if Suppression.query.filter_by(email=contact.email).first():
                contact.status = "suppressed"
                continue
            contact.status = "pending"
            contact.last_error = None
            queued += 1
        db.session.commit()
        flash(f"Queued {queued} failed contact{'s' if queued != 1 else ''} for review. No emails were sent.", "success")
        return redirect(url_for("contacts", status="pending"))

    @app.route("/retry/<int:contact_id>", methods=["POST"])
    @login_required
    def retry_contact(contact_id):
        contact = db.session.get(Contact, contact_id)
        if not contact:
            flash("Contact not found.", "error")
            return redirect(url_for("contacts"))
        if Suppression.query.filter_by(email=contact.email).first():
            contact.status = "suppressed"
            db.session.commit()
            flash("This address is suppressed and cannot be retried.", "error")
            return redirect(url_for("contacts"))
        if contact.status != "failed":
            flash("Only failed contacts can be queued for retry.", "error")
            return redirect(url_for("contacts"))

        contact.status = "pending"
        contact.last_error = None
        db.session.commit()
        flash("Contact queued for retry. Review the preview before sending.", "success")
        return redirect(url_for("contacts"))

    @app.route("/suppress", methods=["POST"])
    @login_required
    def suppress():
        email = request.form.get("email", "").strip().lower()
        try:
            email = validate_email(email, check_deliverability=False).normalized
        except EmailNotValidError:
            flash("Enter a valid email address to suppress.", "error")
            return redirect(url_for("contacts"))
        if not Suppression.query.filter_by(email=email).first():
            db.session.add(Suppression(email=email, reason="manual suppression"))
            c = Contact.query.filter(db.func.lower(Contact.email) == email).first()
            if c:
                c.status = "suppressed"
            db.session.commit()
            flash("Address added to the suppression list.", "success")
        return redirect(url_for("contacts"))

    @app.route("/export")
    @login_required
    def export():
        import csv
        from io import StringIO
        from flask import Response
        query = request.args.get("q", "").strip()
        status = request.args.get("status", "all")
        contacts_query = Contact.query
        if query:
            pattern = f"%{query}%"
            contacts_query = contacts_query.filter(
                or_(Contact.email.ilike(pattern), Contact.company.ilike(pattern),
                    Contact.contact_name.ilike(pattern))
            )
        if status in {"pending", "sent", "failed", "suppressed"}:
            contacts_query = contacts_query.filter_by(status=status)
        s = StringIO()
        w = csv.writer(s)
        w.writerow(["email", "company", "website", "status", "sent_at", "last_error"])
        for c in contacts_query.order_by(Contact.id).all():
            w.writerow([c.email, c.company, c.website, c.status, c.sent_at, c.last_error])
        filename = "outreach_results.csv" if not query and status == "all" else "outreach_contacts_filtered.csv"
        return Response(s.getvalue(), mimetype="text/csv",
                        headers={"Content-Disposition": f"attachment; filename={filename}"})

    return app

if __name__ == "__main__":
    create_app().run()
